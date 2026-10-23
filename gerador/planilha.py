import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

from regras import SLOTS_GRADE, normalizar, turma_pode_usar_slot


CORES_DISCIPLINA = {
    "azul": "9DC3E6",
    "rosa": "F4B6C2",
    "amarelo": "FFD966",
    "verde": "A9D18E",
    "vermelho": "FF9B9B",
    "roxo": "C9B6E4",
    "laranja": "F8CBAD",
    "cinza": "D9D9D9",
    "ciano": "B7DEE8",
    "marrom": "D7B899"
}

verde_titulo = "1B6E1B"
vermelho = "FF0000"
amarelo = "FFF200"
verde_intervalo = "B7F34A"
azul_escuro = "0F2A75"
cinza_vago = "E7E7E7"
laranja_dividida = "F4B183"
laranja_ajuste = "FCE4D6"
vermelho_falta = "F4CCCC"
branco = "FFFFFF"
preto = "000000"

paleta = [
    "9DC3E6",
    "F4B183",
    "FFB7B2",
    "A9D18E",
    "FFD966",
    "B4C7E7",
    "C6E0B4",
    "F4CCCC",
    "D9EAD3",
    "D9D2E9",
    "B7DEE8",
    "F8CBAD",
    "CFE2F3",
    "D5A6BD"
]

borda_fina = Side(style="thin", color="000000")
borda_grossa = Side(style="medium", color="000000")


def nome_arquivo_seguro(nome):
    nome = re.sub(r'[\\/:*?"<>|]', "", nome).strip()
    return nome or "Grade_Automatica"


def borda_simples():
    return Border(
        top=borda_fina,
        bottom=borda_fina,
        left=borda_fina,
        right=borda_fina
    )


def preencher_grade(ws, nome_escola, coordenacao, turmas, agendador, modo, cores_disciplina):
    total_colunas = 2 + len(turmas)
    ultima_coluna = ws.cell(1, total_colunas).column_letter

    # Título
    ws.merge_cells(f"A1:{ultima_coluna}1")
    ws["A1"] = nome_escola
    ws["A1"].font = Font(size=18, bold=True, color=verde_titulo)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A2:{ultima_coluna}2")
    ws["A2"] = datetime.now().strftime("%d/%m/%Y")
    if coordenacao:
        ws["A2"] = f"{ws['A2'].value} - Coordenação: {coordenacao}"
    ws["A2"].font = Font(size=14, bold=True, color=vermelho)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A3:{ultima_coluna}3")
    titulo_modo = (
        "GRADE ESCOLAR AUTOMÁTICA"
        if modo == "normal"
        else "GRADE ESCOLAR - AJUSTES SUGERIDOS"
    )
    ws["A3"] = titulo_modo
    ws["A3"].font = Font(size=15, bold=True, color=preto)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    # Cabeçalho das disciplinas cadastradas
    ws.cell(4, 1).value = "Disciplinas"
    ws.cell(4, 2).value = "Carga/semana"

    for indice, turma in enumerate(turmas, start=3):
        lista_disciplinas = []
        carga_total = 0

        for dados in turma["disciplinas"].values():
            lista_disciplinas.append(
                f"{dados['nome']} ({dados['qtd']}x)"
            )
            carga_total += dados["qtd"]

        ws.cell(4, indice).value = "\n".join(lista_disciplinas)
        ws.cell(4, indice).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Cabeçalho principal
    ws.cell(5, 1).value = "Dia"
    ws.cell(5, 2).value = "Horário"

    for indice, turma in enumerate(turmas, start=3):
        ws.cell(5, indice).value = turma["nome"]

    for linha in (4, 5):
        for cell in ws[linha]:
            if cell.column <= total_colunas:
                cell.fill = PatternFill("solid", fgColor=amarelo)
                cell.font = Font(bold=True, color=preto, size=10)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )
                cell.border = Border(
                    top=borda_grossa,
                    bottom=borda_grossa,
                    left=borda_fina,
                    right=borda_fina
                )

    # Legenda rápida
    legenda_linha = 6
    ws.cell(legenda_linha, 1).value = "LEGENDA"
    ws.cell(legenda_linha, 1).font = Font(bold=True)
    ws.cell(legenda_linha, 2).value = (
        "Aulas = professor + disciplina | "
        "COMP VAGO = sem aula | "
        "AJUSTE = fora do horário exato informado"
    )
    ws.merge_cells(
        start_row=legenda_linha,
        start_column=2,
        end_row=legenda_linha,
        end_column=max(2, total_colunas)
    )
    ws.cell(legenda_linha, 2).alignment = Alignment(wrap_text=True)

    linha = 7
    indice_turma_coluna = {
        turma["chave"]: indice
        for indice, turma in enumerate(turmas, start=3)
    }

    cor_disciplina = {}

    def obter_cor(disciplina, nome):
        if cores_disciplina.get(disciplina):
            return cores_disciplina[disciplina]

        chave = normalizar(nome)

        if chave not in cor_disciplina:
            cor_disciplina[chave] = paleta[
                len(cor_disciplina) % len(paleta)
            ]

        return cor_disciplina[chave]

    agenda = agendador.atribuicoes
    faltas = agendador.contar_faltas()

    for dia in agendador.dias:
        inicio_dia = linha

        for horario, tipo, turno in SLOTS_GRADE:
            ws.cell(linha, 1).value = dia
            ws.cell(linha, 2).value = horario

            ws.cell(linha, 1).font = Font(bold=True)
            ws.cell(linha, 2).font = Font(bold=True)
            ws.cell(linha, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(linha, 2).alignment = Alignment(horizontal="center", vertical="center")

            if tipo == "intervalo":
                ws.cell(linha, 2).fill = PatternFill(
                    "solid",
                    fgColor=verde_intervalo
                )

                texto_intervalo = "INTERVALO"

                for col in range(3, total_colunas + 1):
                    celula = ws.cell(linha, col)
                    celula.value = texto_intervalo
                    celula.fill = PatternFill(
                        "solid",
                        fgColor=azul_escuro
                    )
                    celula.font = Font(bold=True, color=branco)
                    celula.alignment = Alignment(horizontal="center", vertical="center")
                    celula.border = borda_simples()

            else:
                for turma in turmas:
                    col = indice_turma_coluna[turma["chave"]]
                    celula = ws.cell(linha, col)

                    ativa = turma_pode_usar_slot(turma, horario)
                    registro = agenda.get(
                        turma["chave"],
                        {}
                    ).get((dia, horario))

                    if not ativa:
                        celula.value = ""
                        celula.fill = PatternFill("solid", fgColor="FFFFFF")

                    elif not registro:
                        faltas_turma = faltas.get(
                            turma["chave"],
                            {}
                        )

                        if modo == "sugestao" and any(valor > 0 for valor in faltas_turma.values()):
                            pendencias = [
                                dados["nome"]
                                for disciplina, dados in turma["disciplinas"].items()
                                if faltas_turma.get(disciplina, 0) > 0
                            ]

                            celula.value = "COMP VAGO\nFALTAS: " + ", ".join(pendencias)
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=vermelho_falta
                            )
                        else:
                            celula.value = "COMP VAGO"
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=cinza_vago
                            )

                        celula.font = Font(bold=True, color=preto, size=8)

                    elif registro["tipo"] == "normal":
                        texto = (
                            f"{registro['professor']}\n"
                            f"{registro['disciplina_nome']}"
                        )

                        if registro.get("aula_dupla"):
                            texto += "\nAULA DUPLA"

                        if registro.get("ajuste"):
                            texto += "\nAJUSTE"
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=laranja_ajuste
                            )
                        else:
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=obter_cor(
                                    registro["disciplina"],
                                    registro["disciplina_nome"]
                                )
                            )

                        celula.value = texto
                        celula.font = Font(bold=True, color=preto, size=9)

                    else:
                        texto = (
                            f"{registro['professor_a']} - "
                            f"{registro['disciplina_a_nome']}\n"
                            f"{registro['professor_b']} - "
                            f"{registro['disciplina_b_nome']}\n"
                            "AULA DIVIDIDA"
                        )

                        if registro.get("ajuste"):
                            texto += "\nAJUSTE"
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=laranja_ajuste
                            )
                        else:
                            celula.fill = PatternFill(
                                "solid",
                                fgColor=laranja_dividida
                            )

                        celula.value = texto
                        celula.font = Font(bold=True, color=preto, size=8)

                    celula.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True
                    )
                    celula.border = borda_simples()

            for col in range(1, total_colunas + 1):
                ws.cell(linha, col).border = borda_simples()

            linha += 1

        fim_dia = linha - 1

        ws.cell(inicio_dia, 1).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        for row in range(inicio_dia, fim_dia + 1):
            ws.cell(row, 1).fill = PatternFill(
                "solid",
                fgColor="D9EAD3"
            )

        if fim_dia > inicio_dia:
            ws.merge_cells(
                start_row=inicio_dia,
                start_column=1,
                end_row=fim_dia,
                end_column=1
            )

    # Pendências detalhadas abaixo da grade.
    linha += 2

    existe_falta = any(
        quantidade > 0
        for faltas_turma in faltas.values()
        for quantidade in faltas_turma.values()
    )

    if existe_falta:
        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=total_colunas
        )
        ws.cell(linha, 1).value = "DISCIPLINAS QUE NÃO FORAM TOTALMENTE ENCAIXADAS"
        ws.cell(linha, 1).font = Font(
            bold=True,
            color=vermelho,
            size=12
        )
        ws.cell(linha, 1).alignment = Alignment(horizontal="center")
        linha += 1

        for turma in turmas:
            faltas_turma = faltas[turma["chave"]]
            itens = []

            for disciplina, quantidade in faltas_turma.items():
                if quantidade > 0:
                    itens.append(
                        f"{turma['disciplinas'][disciplina]['nome']}: {quantidade} aula(s)"
                    )

            if itens:
                ws.cell(linha, 1).value = turma["nome"]
                ws.cell(linha, 1).font = Font(bold=True)
                ws.cell(linha, 2).value = " | ".join(itens)
                ws.merge_cells(
                    start_row=linha,
                    start_column=2,
                    end_row=linha,
                    end_column=max(2, total_colunas)
                )
                linha += 1

    # Larguras / impressão
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 20

    for col in range(3, total_colunas + 1):
        ws.column_dimensions[ws.cell(5, col).column_letter].width = 27

    for row in range(7, linha):
        ws.row_dimensions[row].height = 56

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 25
    ws.row_dimensions[4].height = 70
    ws.row_dimensions[5].height = 32

    ws.freeze_panes = "C7"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"


def preencher_tabela(ws, titulo, cabecalho, linhas, larguras):
    total_colunas = len(cabecalho)
    ultima_coluna = ws.cell(1, total_colunas).column_letter

    ws.merge_cells(f"A1:{ultima_coluna}1")
    ws["A1"] = titulo
    ws["A1"].font = Font(size=15, bold=True, color=vermelho)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for col, valor in enumerate(cabecalho, start=1):
        celula = ws.cell(2, col)
        celula.value = valor
        celula.fill = PatternFill("solid", fgColor=amarelo)
        celula.font = Font(bold=True, color=preto)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = Border(
            top=borda_grossa,
            bottom=borda_grossa,
            left=borda_fina,
            right=borda_fina
        )

    for linha, valores in enumerate(linhas, start=3):
        for col, valor in enumerate(valores, start=1):
            celula = ws.cell(linha, col)
            celula.value = valor
            celula.alignment = Alignment(vertical="center", wrap_text=True)
            celula.border = borda_simples()

    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(2, col).column_letter].width = largura

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def preencher_erros(ws, erros):
    preencher_tabela(
        ws,
        "O QUE IMPEDIU A GRADE DE FICAR COMPLETA",
        ["Turma", "Disciplina", "Aulas faltando", "Motivo", "Como resolver"],
        [
            [erro["turma"], erro["disciplina"], erro["faltando"], erro["motivo"], erro["solucao"]]
            for erro in erros
        ],
        [28, 24, 15, 60, 50]
    )

    for linha in range(3, len(erros) + 3):
        ws.cell(linha, 3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(linha, 3).fill = PatternFill("solid", fgColor=vermelho_falta)


def preencher_ajustes(ws, ajustes):
    preencher_tabela(
        ws,
        "HORÁRIOS QUE PRECISAM SER LIBERADOS PELOS PROFESSORES",
        ["Professor", "Dia", "Horário", "Turma", "Disciplina"],
        [
            [ajuste["professor"], ajuste["dia"], ajuste["horario"], ajuste["turma"], ajuste["disciplina"]]
            for ajuste in ajustes
        ],
        [26, 18, 18, 28, 24]
    )

    for linha in range(3, len(ajustes) + 3):
        for col in range(1, 6):
            ws.cell(linha, col).fill = PatternFill("solid", fgColor=laranja_ajuste)


def gerar_planilha(nome_escola, coordenacao, turmas, resultado, cores_disciplina):
    wb = Workbook()
    ws = wb.active
    ws.title = "Grade"

    preencher_grade(
        ws,
        nome_escola,
        coordenacao,
        turmas,
        resultado["agendador"],
        "normal",
        cores_disciplina
    )

    if resultado["erros"]:
        preencher_erros(wb.create_sheet("Erros"), resultado["erros"])

    sugestao = resultado["sugestao"]

    if sugestao:
        preencher_grade(
            wb.create_sheet("Solução sugerida"),
            nome_escola,
            coordenacao,
            turmas,
            sugestao,
            "sugestao",
            cores_disciplina
        )
        preencher_ajustes(wb.create_sheet("Ajustes da solução"), sugestao.listar_ajustes())

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)

    return arquivo


def cor_da_disciplina(nome_cor):
    if not nome_cor:
        return None

    return CORES_DISCIPLINA.get(normalizar(nome_cor))


def nome_do_arquivo(nome_escola):
    return nome_arquivo_seguro(f"Grade_{nome_escola}").replace(" ", "_") + ".xlsx"
