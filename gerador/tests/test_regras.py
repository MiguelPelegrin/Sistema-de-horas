import os
import sys
import unittest
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

import regras
from planilha import gerar_planilha


def professor(id, nome, disciplinas, disponibilidade, carga_maxima=None):
    return {
        "id": id,
        "nome": nome,
        "disciplinas": {disciplina: disciplina.title() for disciplina in disciplinas},
        "disponibilidade": {dia: set(horarios) for dia, horarios in disponibilidade.items()},
        "carga_maxima": carga_maxima
    }


def turma(chave, turnos, disciplinas):
    return {
        "id": chave,
        "nome": f"Turma {chave}",
        "chave": chave,
        "turnos": turnos,
        "disciplinas": {
            disciplina: {"nome": disciplina.title(), "qtd": quantidade}
            for disciplina, quantidade in disciplinas.items()
        }
    }


def semana_toda(turno):
    return {dia: regras.SLOTS_TURNO[turno] for dia in regras.SEMANA}


def aulas_do_professor(agendador, professor_id):
    return sorted(
        agendador.agendas_professores[professor_id].keys(),
        key=lambda chave: (regras.SEMANA.index(chave[0]), regras.ORDEM_GRADE.index(chave[1]))
    )


class TestRegras(unittest.TestCase):

    def test_grade_completa_quando_tudo_cabe(self):
        professores = [
            professor(1, "Ana", ["mat"], semana_toda("Manhã")),
            professor(2, "Bruno", ["port"], semana_toda("Manhã"))
        ]
        turmas = [turma("1", ["Manhã"], {"mat": 5, "port": 4})]

        resultado = regras.montar_grade(professores, turmas, [])

        self.assertTrue(resultado["completo"])
        self.assertEqual(resultado["agendador"].total_faltas(), 0)
        self.assertEqual(resultado["erros"], [])

    def test_professor_nunca_fica_em_duas_turmas_ao_mesmo_tempo(self):
        professores = [professor(1, "Ana", ["mat"], semana_toda("Manhã"))]
        turmas = [
            turma("1", ["Manhã"], {"mat": 5}),
            turma("2", ["Manhã"], {"mat": 5})
        ]

        agendador = regras.montar_grade(professores, turmas, [])["agendador"]
        horarios = [
            chave
            for turma_chave in ("1", "2")
            for chave in agendador.atribuicoes[turma_chave]
        ]

        self.assertEqual(len(horarios), 10)
        self.assertEqual(len(horarios), len(set(horarios)))

    def test_intersticio_de_11_horas_entre_noite_e_manha(self):
        disponibilidade = {
            "Segunda-feira": regras.SLOTS_TURNO["Noite"],
            "Terça-feira": regras.SLOTS_TURNO["Manhã"]
        }
        professores = [professor(1, "Junior", ["mat"], disponibilidade)]
        turmas = [
            turma("noite", ["Noite"], {"mat": 2}),
            turma("manha", ["Manhã"], {"mat": 2})
        ]

        agendador = regras.montar_grade(professores, turmas, [])["agendador"]
        aulas = aulas_do_professor(agendador, 1)
        segunda = [horario for dia, horario in aulas if dia == "Segunda-feira"]
        terca = [horario for dia, horario in aulas if dia == "Terça-feira"]

        self.assertTrue(segunda and terca)
        descanso = 24 * 60 - regras.fim_minutos(max(segunda, key=regras.fim_minutos)) \
            + regras.inicio_minutos(min(terca, key=regras.inicio_minutos))
        self.assertGreaterEqual(descanso, regras.INTERSTICIO)

    def test_carga_horaria_maxima_e_respeitada(self):
        professores = [professor(1, "Eva", ["fis"], semana_toda("Manhã"), carga_maxima=3)]
        turmas = [turma("1", ["Manhã"], {"fis": 5})]

        resultado = regras.montar_grade(professores, turmas, [])

        self.assertEqual(len(resultado["agendador"].agendas_professores[1]), 3)
        self.assertEqual(resultado["agendador"].total_faltas(), 2)
        self.assertEqual(resultado["erros"][0]["tipo"], "carga")

    def test_aula_dupla_nao_atravessa_o_intervalo(self):
        self.assertTrue(regras.slot_adjacente("07h20 - 08h10", "08h10 - 09h00"))
        self.assertFalse(regras.slot_adjacente("09h00 - 09h50", "10h05 - 10h55"))
        self.assertFalse(regras.slot_adjacente("14h20 - 15h10", "18h10 - 19h00"))

    def test_dias_letivos_limitam_a_grade(self):
        professores = [professor(1, "Ana", ["mat"], semana_toda("Manhã"))]
        turmas = [turma("1", ["Manhã"], {"mat": 4})]

        agendador = regras.montar_grade(professores, turmas, [], dias=["Segunda-feira", "Terça-feira"])["agendador"]
        dias_usados = {dia for dia, _ in agendador.atribuicoes["1"]}

        self.assertTrue(dias_usados <= {"Segunda-feira", "Terça-feira"})

    def test_disciplina_sem_professor_vira_erro(self):
        professores = [professor(1, "Ana", ["mat"], semana_toda("Manhã"))]
        turmas = [turma("1", ["Manhã"], {"mat": 2, "ingles": 2})]

        resultado = regras.montar_grade(professores, turmas, [])

        self.assertFalse(resultado["completo"])
        self.assertEqual(len(resultado["erros"]), 1)
        self.assertEqual(resultado["erros"][0]["tipo"], "sem_professor")
        self.assertIsNone(resultado["sugestao"])

    def test_solucao_sugerida_quando_falta_disponibilidade(self):
        professores = [professor(1, "Andreia", ["port"], {"Segunda-feira": ["07h20 - 08h10"]})]
        turmas = [turma("1", ["Manhã"], {"port": 3})]

        resultado = regras.montar_grade(professores, turmas, [])

        self.assertEqual(resultado["erros"][0]["tipo"], "pouca_disponibilidade")
        self.assertIsNotNone(resultado["sugestao"])
        self.assertLess(resultado["sugestao"].total_faltas(), resultado["agendador"].total_faltas())
        self.assertTrue(resultado["sugestao"].listar_ajustes())

    def test_aula_dividida_usa_dois_professores_no_mesmo_horario(self):
        dois_horarios = {"Segunda-feira": ["07h20 - 08h10", "08h10 - 09h00"]}
        professores = [
            professor(1, "Carla", ["web"], dois_horarios),
            professor(2, "Davi", ["bd"], dois_horarios)
        ]
        turmas = [turma("1", ["Manhã"], {"web": 2, "bd": 2})]
        divisoes = [{"turma": "1", "disciplina_a": "web", "disciplina_b": "bd"}]

        agendador = regras.montar_grade(professores, turmas, divisoes)["agendador"]
        tipos = [registro["tipo"] for registro in agendador.atribuicoes["1"].values()]

        self.assertEqual(agendador.total_faltas(), 0)
        self.assertEqual(tipos, ["dividida", "dividida"])

    def test_resultado_e_sempre_o_mesmo(self):
        professores = [professor(i, f"P{i}", ["mat", "port"], semana_toda("Manhã")) for i in range(1, 4)]
        turmas = [turma(str(i), ["Manhã"], {"mat": 4, "port": 4}) for i in range(1, 5)]

        primeira = regras.montar_grade(professores, turmas, [])["agendador"].atribuicoes
        segunda = regras.montar_grade(professores, turmas, [])["agendador"].atribuicoes

        self.assertEqual(primeira, segunda)


class TestPlanilha(unittest.TestCase):

    def test_planilha_com_erros_e_solucao_tem_todas_as_abas(self):
        professores = [professor(1, "Andreia", ["port"], {"Segunda-feira": ["07h20 - 08h10"]})]
        turmas = [turma("1", ["Manhã"], {"port": 3})]
        resultado = regras.montar_grade(professores, turmas, [])

        arquivo = gerar_planilha("ETEC Teste", "Coordenação", turmas, resultado, {"port": "9DC3E6"})
        planilha = load_workbook(BytesIO(arquivo.getvalue()))

        self.assertEqual(planilha.sheetnames, ["Grade", "Erros", "Solução sugerida", "Ajustes da solução"])
        self.assertEqual(planilha["Grade"]["A1"].value, "ETEC Teste")
        self.assertEqual(planilha["Erros"].cell(3, 2).value, "Port")

    def test_planilha_completa_tem_so_a_grade(self):
        professores = [professor(1, "Ana", ["mat"], semana_toda("Manhã"))]
        turmas = [turma("1", ["Manhã"], {"mat": 2})]
        resultado = regras.montar_grade(professores, turmas, [])

        planilha = load_workbook(BytesIO(gerar_planilha("ETEC", None, turmas, resultado, {}).getvalue()))

        self.assertEqual(planilha.sheetnames, ["Grade"])


if __name__ == "__main__":
    unittest.main()
